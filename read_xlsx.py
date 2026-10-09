# -*- coding: utf-8 -*-
import openpyxl, json, sys
wb = openpyxl.load_workbook('待下载文献20260926.xlsx', data_only=True)
out = []
for ws in wb.worksheets:
    out.append(f'=== SHEET: {ws.title} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}')
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 8:
            out.append(f'ROW{i+1}: {row}')
    # header detection
    hdr = [c.value for c in ws[1]]
    out.append('HEADER: ' + json.dumps(hdr, ensure_ascii=False))
with open('xlsx_preview.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('sheets:', [ws.title for ws in wb.worksheets])
