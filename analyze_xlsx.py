# -*- coding: utf-8 -*-
import openpyxl, json, collections
wb = openpyxl.load_workbook('待下载文献20260926.xlsx', data_only=True)
report = []
for ws in wb.worksheets:
    report.append(f'=== SHEET: {ws.title} max_row={ws.max_row} max_col={ws.max_column}')
    hdr = [str(c.value) if c.value else '' for c in ws[1]]
    report.append('HEADER: ' + json.dumps(hdr, ensure_ascii=False))
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    report.append(f'data rows: {len(rows)}')
    # find DOI and title column indices
    doi_i = next((i for i, h in enumerate(hdr) if 'DOI' in h.upper()), None)
    title_i = next((i for i, h in enumerate(hdr) if '英文' in h and '题目' in h), None)
    src_i = next((i for i, h in enumerate(hdr) if h == '来源'), None)
    link_i = next((i for i, h in enumerate(hdr) if '链接' in h), None)
    n_doi = sum(1 for r in rows if doi_i is not None and r[doi_i])
    n_title = sum(1 for r in rows if title_i is not None and r[title_i])
    n_link = sum(1 for r in rows if link_i is not None and r[link_i])
    report.append(f'with DOI: {n_doi}, with title: {n_title}, with link: {n_link}')
    if src_i is not None:
        srcs = collections.Counter(str(r[src_i]) for r in rows)
        report.append('sources: ' + json.dumps(dict(srcs), ensure_ascii=False))
    if ws.title != 'Sheet1':
        for i, r in enumerate(rows[:5]):
            report.append(f'ROW{i+2}: {json.dumps([str(x)[:60] if x else "" for x in r], ensure_ascii=False)}')

with open('xlsx_report.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(report))
print('ok')
